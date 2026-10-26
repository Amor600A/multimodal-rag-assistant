# -*- coding: utf-8 -*-
"""创建 Milvus Lite 知识库脚本。

流程：读取 PDF/DOCX/TXT -> 提取文本 -> 切片 -> 嵌入 -> 写入 Milvus Lite；
图片/视频不切片，直接多模态嵌入写单条记录。
数据文件默认落在 data/milvus/rag_knowledge.db（由配置 MILVUS_DB_PATH 决定）。

支持格式：
    .pdf  -> 逐页提取，页码精确到页
    .docx -> 无固定分页（分页由渲染决定），用切片序号定位（source 记为 文件名#c块序号）
    .txt  -> 纯文本，UTF-8 优先、GB18030 兜底，无分页，同样用切片序号定位
    .jpg/.jpeg/.png/.bmp -> 图片不切分，一个文件一条记录（需 EMBEDDING_BACKEND=dashscope_mm）
    .mp4/.mov/.avi/.webm（公网 URL）-> 视频不切分，一条记录，同样需要多模态嵌入

图片/视频记录没有正文，text 字段存 --description（缺省用文件名），
用于检索结果展示与关键词过滤。所有记录都写入 file_type 字段（pdf/docx/txt/png/video），
便于按来源类型区分；docx/txt 因无固定分页，source 用 #c块序号 而非 #p页码。

改造说明（相对最初参考脚本）：
1. 嵌入改用项目统一嵌入器（EMBEDDING_BACKEND=dashscope_mm 时为
   DashScope multimodal-embedding-v1，文本/图片/视频同处一个 1024 维语义空间），
   与线上检索共用同一向量空间
2. 存储改用项目 MilvusLite（接口与 MockMilvus 一致），不再依赖 langchain_milvus
3. 每条切片补充权限标量字段（scope/visibility/owner_user_id/dept_id）与
   来源类型字段 file_type，否则检索时的权限 expr 无法命中、也无法按类型区分
4. 定位信息写入 source 字段：PDF 为 文件名#p页码，docx/txt 为 文件名#c块序号，
   不额外占用 schema 字段

用法：
    python -m app.infra.vectordb.create_knowledge_base 文档.pdf
    python -m app.infra.vectordb.create_knowledge_base 文档.docx --visibility dept --dept-id <部门ID>
    python -m app.infra.vectordb.create_knowledge_base 手册.txt
    python -m app.infra.vectordb.create_knowledge_base 文档.pdf --visibility private --owner-user-id <用户ID>
    python -m app.infra.vectordb.create_knowledge_base 图片.png --description "园区导览图"
    python -m app.infra.vectordb.create_knowledge_base https://example.com/a.mp4 --description "入园须知视频"

批量入库（目录约定，权限标记由目录结构决定，无需逐个指定 --visibility）：
    python -m app.infra.vectordb.create_knowledge_base --dir data/knowledge
    目录结构（相对 --dir 根目录）：
        public/...              -> visibility=public
        dept/<dept_id>/...      -> visibility=dept，dept_id 取自目录名
        private/<user_id>/...   -> visibility=private，owner_user_id 取自目录名
        根目录及其他目录下的文件  -> visibility=public
        各目录下 videos.txt     -> 视频清单（VIDEO_KNOWLEDGE = [{"url":..., "description":...}]），
                                   权限同所在目录
    批量模式扫描本地文件（.pdf/.docx/.txt/图片）与各目录的 videos.txt 清单；
    单文件/单条 URL 失败不中断整批，结束打印成功/失败/跳过汇总。

注意：Milvus Lite 为单文件独占，运行本脚本前请先停止后端服务，避免文件被占用。
      .txt 需为 UTF-8 或 GBK 编码。
"""

from __future__ import annotations

import argparse
import ast
import hashlib
import os
import re
from typing import List, Tuple
from pypdf import PdfReader

from app.core.config import get_settings
from app.infra.vectordb.milvus_lite import MilvusLite

# 图片格式：本地文件，走 Base64 上传
_IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".bmp"}
# 视频格式：仅支持公网 URL
_VIDEO_EXTS = {".mp4", ".mov", ".avi", ".webm"}

# 批量入库支持的本地文件类型；视频以各目录 videos.txt 清单中的公网 URL 参与，不走文件扫描
_BATCH_EXTS = {".pdf", ".docx", ".txt"} | _IMAGE_EXTS

# 权限标量白名单：与权限表达式构造的校验规则一致（仅字母/数字/下划线/中划线），
# dept_id/owner_user_id 在入库时 fail-fast，避免写入后检索期才暴露
_ID_RE = re.compile(r"^[A-Za-z0-9_-]+$")

# 默认切片参数（与参考脚本保持一致）
CHUNK_SIZE = 500
CHUNK_OVERLAP = 100

# 分隔符优先级：段落 -> 换行 -> 中文句末 -> 英文句点 -> 空格
_SEPARATORS = ("\n\n", "\n", "。", "！", "？", ".", " ")


def extract_text_from_pdf(pdf) -> Tuple[str, List[int]]:
    """从 PDF 提取文本，并记录每个字符对应的页码。

    参数:
        pdf: PdfReader 对象
    返回:
        text: 提取的文本内容
        char_page_mapping: 与 text 等长的页码列表
    """
    text = ""
    char_page_mapping: list[int] = []

    for page_number, page in enumerate(pdf.pages, start=1):
        extracted_text = page.extract_text()
        if extracted_text:
            text += extracted_text
            # 为当前页面的每个字符记录页码
            char_page_mapping.extend([page_number] * len(extracted_text))
        else:
            print(f"第 {page_number} 页未提取到文本，已跳过。")

    return text, char_page_mapping

def extract_text_from_docx(docx_path: str) -> Tuple[str, List[int]]:
    """从 DOCX 提取文本。

    DOCX 没有固定分页（分页由渲染时决定），无法逐页定位，
    因此页码占位统一记为 1；source 定位不使用页码，
    由 build_text_chunks 以切片序号记为 文件名#c块序号。

    提取范围：正文段落 + 表格单元格文本（攻略/政策类文档常用表格排版）。

    参数:
        docx_path: DOCX 文件路径
    返回:
        text: 提取的文本内容
        char_page_mapping: 与 text 等长的页码列表（全部为 1）
    """
    from docx import Document  # 延迟导入：仅处理 DOCX 时要求 python-docx

    document = Document(docx_path)
    parts: list[str] = []
    for para in document.paragraphs:
        t = para.text.strip()
        if t:
            parts.append(t)
    for table in document.tables:
        for row in table.rows:
            # 横向合并单元格（gridSpan）会让 row.cells 把同一个单元格返回多次，
            # 按底层 XML 节点去重；不能用文本去重，因为两列内容相同是合法的
            cells: list[str] = []
            seen: set = set()
            for cell in row.cells:
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)
                t = cell.text.strip()
                if t:
                    cells.append(t)
            line = " | ".join(cells)
            if line:
                parts.append(line)

    text = "\n".join(parts)
    return text, [1] * len(text)

def extract_text_from_txt(txt_path: str) -> Tuple[str, List[int]]:
    """从纯文本文件提取文本。

    编码顺序：UTF-8（含 BOM）优先；中文 Windows 下记事本常另存为 GBK，
    失败时按 GB18030 解码（GBK 的超集，不会静默产出乱码）。

    参数:
        txt_path: 文本文件路径
    返回:
        text: 提取的文本内容
        char_page_mapping: 与 text 等长的页码列表（全部为 1）
    """
    with open(txt_path, "rb") as f:
        raw = f.read()

    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("gb18030")

    return text, [1] * len(text)

def _media_embedder():
    """取出多模态嵌入器；图片/视频入库必须使用统一多模态空间。"""
    # 嵌入器保持延迟导入：只有处理图片/视频时才需要加载 dashscope 依赖
    from app.rag.embedder import MultimodalEmbedder, get_embedder

    embedder = get_embedder()
    if not isinstance(embedder, MultimodalEmbedder):
        raise SystemExit(
            "图片/视频入库需要 EMBEDDING_BACKEND=dashscope_mm，"
            f"当前为 {get_settings().embedding_backend}"
        )
    return embedder

def embed_image(image_path: str, embedder=None):
    """生成图片向量（1024 维，与文本同一语义空间）。

    由嵌入器负责读取本地文件并转 Base64，支持 jpg/jpeg/png/bmp。
    批量模式下可外部注入 embedder 复用单例，避免重复取。
    """
    if embedder is None:
        embedder = _media_embedder()
    return embedder.embed_image(image_path)

def embed_video(video_url: str, embedder=None):
    """生成视频向量（1024 维，与文本同一语义空间）。

    接口只接受公网可访问的 http(s) URL，本地视频文件需先上传到可访问地址。
    单条视频返回一个向量，不做多帧平均。
    批量模式下可外部注入 embedder 复用单例，避免重复取。
    """
    if embedder is None:
        embedder = _media_embedder()
    return embedder.embed_video(video_url)

def _split_with_offsets(
    text: str, chunk_size: int = CHUNK_SIZE, chunk_overlap: int = CHUNK_OVERLAP
) -> List[Tuple[str, int]]:
    """按语义边界切片，并返回每块在原文中的起始偏移。

    与直接按长度硬切相比，偏移量让页码映射可以精确定位到原文位置，
    不会因 chunk_overlap 导致字符位置漂移。

    参数:
        text: 待切片文本
        chunk_size: 单块最大字符数（默认 500）
        chunk_overlap: 块间重叠字符数（默认 100）
    返回:
        [(chunk_text, start_offset), ...]
    """
    chunks: list[tuple[str, int]] = []
    # 全空白文本提前返回空列表，避免无意义循环
    if not text.strip():
        return chunks

    total = len(text)
    start = 0

    while start < total:
        end = min(start + chunk_size, total)
        if end < total:
            # 只在窗口后半段找边界，避免切出过短的块
            window = text[start:end]
            cut = -1
            for sep in _SEPARATORS:
                idx = window.rfind(sep)
                if idx != -1 and idx >= chunk_size // 2:
                    cut = idx + len(sep)
                    break
            if cut > 0:
                end = start + cut

        chunk = text[start:end].strip()
        if chunk:
            chunks.append((chunk, start))

        if end >= total:
            break
        # 向前回退 overlap 个字符，保持块间上下文连续
        start = max(end - chunk_overlap, start + 1)

    print(f"文本被分割成 {len(chunks)} 个块。")
    return chunks

def compute_chunk_pages(
    chunks: List[Tuple[str, int]],
    char_page_mapping: List[int],
    verbose: bool = False,
) -> dict:
    """为每个文本块计算主要来源页码（取块内字符页码的众数）。

    参数:
        chunks: [(chunk_text, start_offset), ...]
        char_page_mapping: 每个字符对应的页码列表
        verbose: True 时打印每个块的完整文本（调试用）；默认 False 只打印汇总
    返回:
        page_info: {文本块内容: 页码}
    """
    page_info: dict[str, int] = {}
    for i, (chunk, offset) in enumerate(chunks, start=1):
        # 按块在原文中的位置取对应页码范围
        page_slice = char_page_mapping[offset:offset + len(chunk)]
        if page_slice:
            page_counts: dict[int, int] = {}
            for page in page_slice:
                page_counts[page] = page_counts.get(page, 0) + 1
            most_common_page = max(page_counts.items(), key=lambda kv: kv[1])[0]
        else:
            most_common_page = 1

        page_info[chunk] = most_common_page

        if verbose:
            print(f"\n--- 第 {i} 个文本块 (长度: {len(chunk)} 字符, 页码: {most_common_page}) ---")
            print(chunk)

    if verbose:
        print()
    print(f"页码映射完成，共 {len(page_info)} 个文本块")
    return page_info

def _write_records(records: List[dict], save_path: str, summary: str, db=None):
    """把记录写入 Milvus Lite 并打印统计（文本切片与媒体记录共用的写库逻辑）。

    参数:
        records: 待写入的记录列表（含 id/text/标量字段，媒体记录另带 vector）
        save_path: Milvus Lite 数据文件路径
        summary: 写入条数的说明文本（如"切片""媒体记录（不切分）"），仅用于日志
        db: 可选，外部传入的 MilvusLite 实例（批量入库时复用同一连接，
            避免每个文件重复打开同一数据文件）；缺省时按 save_path 新建
    返回:
        已写入的 MilvusLite 实例
    """
    settings = get_settings()
    if db is None:
        db = MilvusLite(
            db_path=save_path,
            collection_name=settings.milvus_collection,
            dim=settings.embedding_dim,
        )

    db.insert(records)
    print(f"已写入 {len(records)} 条{summary}，数据文件: {save_path}")
    print(f"当前集合记录总数: {db.count()}（collection={settings.milvus_collection}）")
    print(f"嵌入后端: {settings.embedding_backend}（model={settings.embedding_model}）")
    return db

def build_text_chunks(
    chunks: List[str],
    page_info: dict,
    doc_name: str,
    save_path: str,
    visibility: str,
    dept_id: str | None,
    owner_user_id: str | None,
    file_type: str = "",
    use_page: bool = True,
    db=None,
):
    """把文本块与权限标量写入 Milvus Lite。

    参数:
        chunks: 文本块列表
        page_info: {文本块内容: 页码}（use_page=False 时不参与 source 拼接）
        doc_name: 来源文档名（写入 source 字段）
        save_path: Milvus Lite 数据文件路径
        visibility: public | dept | private
        dept_id: visibility=dept 时的部门 ID
        owner_user_id: visibility=private 时的归属用户 ID
        file_type: 来源文件类型（pdf/docx/txt），写入 file_type 字段
        use_page: True 时 source 记 文件名#p页码（PDF 有真实分页）；
                  False 时记 文件名#c块序号（docx/txt 无固定分页，用切片序号定位）
        db: 可选，外部传入的 MilvusLite 实例（批量入库时复用同一连接，
            避免每个文件重复打开同一数据文件）；缺省时按 save_path 新建
    返回:
        已写入的 MilvusLite 实例
    """
    # private 归入 personal 区，其余归入 domain 区（与 knowledge 路由的打标规则一致）
    scope = "personal" if visibility == "private" else "domain"

    records: list[dict] = []
    # 主键前缀取文件类型（pdf/docx/txt），缺省时用 doc 兜底，避免出现 "-哈希" 这种无意义 id
    id_prefix = file_type or "doc"
    for i, chunk in enumerate(chunks):
        # PDF 用页码定位；docx/txt 无固定分页，用切片序号（从 1 开始）定位
        locator = f"#p{page_info[chunk]}" if use_page else f"#c{i + 1}"
        # 用内容哈希做 id：同一份 PDF 重复建库时幂等覆盖，不会产生重复块
        digest = hashlib.sha1(f"{doc_name}|{locator}|{chunk}".encode("utf-8")).hexdigest()[:16]
        records.append({
            "id": f"{id_prefix}-{digest}",
            "text": chunk,
            "scope": scope,
            "visibility": visibility,
            "owner_user_id": owner_user_id,
            "dept_id": dept_id,
            "source": f"{doc_name}{locator}",
            "file_type": file_type,
        })

    return _write_records(records, save_path, "切片", db=db)

def build_media(
    vector,
    description: str,
    doc_name: str,
    save_path: str,
    visibility: str,
    dept_id: str | None,
    owner_user_id: str | None,
    id_prefix: str,
    file_type: str = "",
    db=None,
):
    """写入单条媒体记录：图片/视频不切分，一个文件对应一个向量。

    参数:
        vector: 多模态嵌入向量（1024 维）
        description: 写入 text 字段的说明文本（检索结果展示用）
        doc_name: 来源文件名或视频 URL（写入 source 字段）
        save_path: Milvus Lite 数据文件路径
        visibility: public | dept | private
        dept_id: visibility=dept 时的部门 ID
        owner_user_id: visibility=private 时的归属用户 ID
        id_prefix: 主键前缀，图片 img、视频 vid
        file_type: 来源文件类型（png/jpg/jpeg/bmp/video），写入 file_type 字段
        db: 可选，外部传入的 MilvusLite 实例（批量入库时复用同一连接）；
            缺省时按 save_path 新建
    返回:
        已写入的 MilvusLite 实例
    """
    scope = "personal" if visibility == "private" else "domain"
    digest = hashlib.sha1(f"{doc_name}|{description}".encode("utf-8")).hexdigest()[:16]
    records = [{
        "id": f"{id_prefix}-{digest}",
        "vector": vector,
        "text": description,
        "scope": scope,
        "visibility": visibility,
        "owner_user_id": owner_user_id,
        "dept_id": dept_id,
        "source": doc_name,
        "file_type": file_type,
    }]

    return _write_records(records, save_path, "媒体记录（不切分）", db=db)

def process_text_with_splitter(
    text: str,
    char_page_mapping: List[int],
    doc_name: str,
    save_path: str,
    visibility: str = "public",
    dept_id: str | None = None,
    owner_user_id: str | None = None,
    file_type: str = "",
    use_page: bool = True,
    db=None,
    chunk_size: int = CHUNK_SIZE,
    chunk_overlap: int = CHUNK_OVERLAP,
    verbose: bool = False,
):
    """处理文本并创建向量存储（保留原函数名与流程）。"""
    chunks = _split_with_offsets(text, chunk_size=chunk_size, chunk_overlap=chunk_overlap)
    # 切片后仍为空（全空白文档），跳过不写库
    if not chunks:
        print(f"切片后无有效内容，跳过入库：{doc_name}")
        return db
    page_info = compute_chunk_pages(chunks, char_page_mapping, verbose=verbose)
    return build_text_chunks(
        [chunk for chunk, _ in chunks],
        page_info,
        doc_name,
        save_path,
        visibility,
        dept_id,
        owner_user_id,
        file_type,
        use_page,
        db=db,
    )

def process_file(
    file_arg: str,
    visibility: str,
    dept_id: str | None,
    owner_user_id: str | None,
    db_path: str,
    description: str | None = None,
    db=None,
    chunk_size: int = CHUNK_SIZE,
    chunk_overlap: int = CHUNK_OVERLAP,
    verbose: bool = False,
    embedder=None,
) -> None:
    """处理单个文件/视频 URL 入库（从 main 抽出的单文件流程，逻辑保持不变）。

    参数:
        file_arg: 文档/图片本地路径，或视频 http(s) URL
        visibility: public | dept | private
        dept_id: visibility=dept 时的部门 ID
        owner_user_id: visibility=private 时的归属用户 ID
        db_path: Milvus Lite 数据文件路径
        description: 图片/视频说明文本，缺省用文件名
        db: 可选，复用的 MilvusLite 实例（批量模式避免重复打开同一数据文件）
        chunk_size: 切片单块最大字符数（默认 500）
        chunk_overlap: 切片块间重叠字符数（默认 100）
        verbose: True 时打印每个切片的完整文本（调试用）
        embedder: 可选，预取的多模态嵌入器（批量模式复用单例，避免重复取）
    """
    is_url = file_arg.lower().startswith(("http://", "https://"))
    # URL 可能带查询串/签名参数（?token=...），先截断再取扩展名，否则匹配不上视频格式
    path_part = file_arg.split("?", 1)[0].split("#", 1)[0] if is_url else file_arg
    ext = os.path.splitext(path_part)[1].lower()

    # 权限标量入库前做白名单校验：dept_id/owner_user_id 最终要进检索期的权限表达式，
    # 与 expr 构造的白名单规则（_ID_RE）保持一致，非法值在此 fail-fast
    for label, value in (("dept_id", dept_id), ("owner_user_id", owner_user_id)):
        if value is not None and not _ID_RE.fullmatch(value):
            raise ValueError(f"{label} 含非法字符: {value}（仅允许字母、数字、下划线、中划线）")

    if not is_url and not os.path.isfile(file_arg):
        raise SystemExit(f"文件不存在: {file_arg}")

    # 图片：本地文件，一条记录，不切片
    if not is_url and ext in _IMAGE_EXTS:
        doc_name = os.path.basename(file_arg)
        description = description or doc_name
        vector = embed_image(file_arg, embedder=embedder)
        build_media(
            vector, description, doc_name, db_path,
            visibility, dept_id, owner_user_id, "img",
            file_type=ext.lstrip("."), db=db,
        )
        return

    # 视频：仅公网 URL，一条记录，不切片
    if is_url and ext in _VIDEO_EXTS:
        description = description or os.path.basename(path_part)
        vector = embed_video(file_arg, embedder=embedder)
        build_media(
            vector, description, file_arg, db_path,
            visibility, dept_id, owner_user_id, "vid",
            file_type="video", db=db,
        )
        return

    # PDF 有真实分页，按页码定位；docx/txt 无固定分页，按切片序号定位
    use_page = ext == ".pdf"
    if ext == ".pdf":
        pdf_reader = PdfReader(file_arg)
        text, char_page_mapping = extract_text_from_pdf(pdf_reader)
    elif ext == ".docx":
        text, char_page_mapping = extract_text_from_docx(file_arg)
    elif ext == ".txt":
        text, char_page_mapping = extract_text_from_txt(file_arg)
    elif ext == ".doc":
        raise SystemExit(f"暂不支持的文件类型: {ext}，需要先手动另存为 .docx 再入库")
    else:
        raise SystemExit(
            f"暂不支持的文件类型: {ext}（支持 .pdf/.docx/.txt，"
            "图片 .jpg/.jpeg/.png/.bmp，视频以 http(s) URL 传入）"
        )
    print(f"提取的文本长度: {len(text)} 个字符。")
    if not text.strip():
        # 扫描件/图片型文档没有文字层；抛 ValueError 让批量模式按单文件失败处理，不中断整批
        raise ValueError(f"未提取到文本内容：{file_arg}（扫描件需先 OCR 再入库）")

    doc_name = os.path.basename(file_arg)
    process_text_with_splitter(
        text,
        char_page_mapping,
        doc_name,
        db_path,
        visibility,
        dept_id,
        owner_user_id,
        file_type=ext.lstrip("."),
        use_page=use_page,
        db=db,
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        verbose=verbose,
    )

def _resolve_visibility(rel_path: str) -> Tuple[str, str | None, str | None]:
    """按目录约定从相对路径推断权限标记。

    目录约定（相对 --dir 根目录）：
        dept/<dept_id>/...      -> ("dept", <dept_id>, None)
        private/<user_id>/...   -> ("private", None, <user_id>)
        根目录及其他目录下的文件  -> ("public", None, None)
    dept/private 一级目录下缺少 ID 子目录时抛 ValueError，
    避免把本应是部门/个人文档的文件误标为 public 造成越权可见。
    """
    parts = rel_path.replace(os.sep, "/").split("/")
    if parts[0] == "dept":
        if len(parts) < 3:
            raise ValueError(f"目录结构不符合约定：{rel_path}（应为 dept/<dept_id>/<文件>）")
        return "dept", parts[1], None
    if parts[0] == "private":
        if len(parts) < 3:
            raise ValueError(f"目录结构不符合约定：{rel_path}（应为 private/<user_id>/<文件>）")
        return "private", None, parts[1]
    return "public", None, None

def _scope_tag(visibility: str, dept_id: str | None, owner_user_id: str | None) -> str:
    """拼权限标记的日志片段，如 "visibility=dept，dept_id=mkt-dept-001"。"""
    tag = f"visibility={visibility}"
    if dept_id:
        tag += f"，dept_id={dept_id}"
    if owner_user_id:
        tag += f"，owner_user_id={owner_user_id}"
    return tag

def _parse_video_manifest(text: str) -> List[dict]:
    """把 videos.txt 内容解析为视频条目列表。

    清单格式为 Python 字面量（与 VIDEO_KNOWLEDGE 的写法一致）：
        VIDEO_KNOWLEDGE = [
            {"url": "https://...", "description": "可选说明"},
        ]
    用 ast.literal_eval 而非 eval，只允许字面量、不执行任意代码。
    """
    # 跳过变量名与等号，只取 [...] 列表部分
    start = text.find("[")
    end = text.rfind("]")
    if start == -1 or end < start:
        raise ValueError("videos.txt 格式错误：未找到 [...] 条目列表")
    try:
        entries = ast.literal_eval(text[start:end + 1])
    except (ValueError, SyntaxError) as exc:
        raise ValueError(f"videos.txt 解析失败：{exc}") from exc
    if not isinstance(entries, list) or not all(isinstance(e, dict) for e in entries):
        raise ValueError("videos.txt 格式错误：条目列表应为 [ {\"url\": ...}, ... ]")
    return entries


def _process_video_manifest(
    manifest_path: str, rel: str, db_path: str, db, embedder=None
) -> Tuple[List[str], List[Tuple[str, str]]]:
    """处理 videos.txt 视频清单，URL 继承清单所在目录的权限标记。

    单条格式错误（缺 url、非 http(s) URL、非支持的视频扩展名）预先拦截记为失败，
    避免传入 process_file 变成 SystemExit 中止整批。
    返回 (成功条目, 失败条目)，格式与 process_directory 的汇总列表一致。
    embedder 由 process_directory 预取后透传，避免重复初始化。
    """
    try:
        visibility, dept_id, owner_user_id = _resolve_visibility(rel)
        text, _page = extract_text_from_txt(manifest_path)
        entries = _parse_video_manifest(text)
    except (ValueError, OSError) as exc:
        # 清单本身不可用时，其中所有 URL 都无法打标/读取，整体记一条失败
        return [], [(rel, str(exc))]

    successes: list[str] = []
    failures: list[tuple[str, str]] = []
    for idx, entry in enumerate(entries, start=1):
        url = str(entry.get("url", "")).strip()
        desc = entry.get("description")
        description = str(desc).strip() if desc else None
        item = f"{rel} 第{idx}条 {url or entry}"
        # URL 可能带查询串/签名参数，先截断再取扩展名（与 process_file 规则一致）
        path_part = url.split("?", 1)[0].split("#", 1)[0]
        ext = os.path.splitext(path_part)[1].lower()
        if not url.lower().startswith(("http://", "https://")) or ext not in _VIDEO_EXTS:
            failures.append((f"{rel} 第{idx}条", f"不是支持的视频 URL：{url or entry}"))
            continue
        try:
            print(f"\n>>> 入库视频 {item}（{_scope_tag(visibility, dept_id, owner_user_id)}）")
            process_file(url, visibility, dept_id, owner_user_id, db_path,
                         description=description, db=db, embedder=embedder)
        except Exception as exc:
            failures.append((item, str(exc)))
            print(f"入库失败：{item} -> {exc}")
            continue
        successes.append(item)
    return successes, failures

def process_directory(
    dir_path: str,
    db_path: str,
    chunk_size: int = CHUNK_SIZE,
    chunk_overlap: int = CHUNK_OVERLAP,
    verbose: bool = False,
) -> None:
    """按目录约定批量入库：遍历 --dir 下所有支持的文件，按所在目录打权限标。

    目录下的 videos.txt 视为视频 URL 清单逐行入库（不作为普通文本参与切片），
    URL 继承所在目录的权限标记。单个文件/单条 URL 失败（如嵌入接口报错、
    目录结构不符约定）不中断整批，结束时打印成功/失败/跳过汇总。
    SystemExit 类环境错误（如嵌入后端不支持图片/视频）对所有同类文件必然重复发生，
    直接中止整批，不逐个重试。

    参数:
        dir_path: 批量入库根目录
        db_path: Milvus Lite 数据文件路径
        chunk_size: 切片单块最大字符数（默认 500）
        chunk_overlap: 切片块间重叠字符数（默认 100）
        verbose: True 时打印每个切片的完整文本（调试用）
    """
    if not os.path.isdir(dir_path):
        # 提前检查：避免静默产出"成功 0 个"，也避免为不存在的目录创建空数据文件
        raise SystemExit(f"批量入库目录不存在: {dir_path}")

    settings = get_settings()
    # 整批复用同一实例：避免为每个文件重复打开同一个 Milvus Lite 数据文件
    db = MilvusLite(
        db_path=db_path,
        collection_name=settings.milvus_collection,
        dim=settings.embedding_dim,
    )

    # dashscope_mm 模式下预取一次多模态嵌入器，后续 process_file/_process_video_manifest
    # 透传进去，避免每个媒体文件都重走一次 get_embedder() + 类型校验
    embedder = None
    if settings.embedding_backend == "dashscope_mm":
        try:
            embedder = _media_embedder()
        except SystemExit as exc:
            # 嵌入器不可用（如 key 未配）直接中止整批，所有媒体文件必然重复失败
            raise SystemExit(f"嵌入器初始化失败：{exc}") from exc

    # 预扫统计总文件数，用于进度条显示
    all_items: list[tuple[str, str, str]] = []  # (file_path, rel, kind)，kind=file|manifest
    for root, _dirs, files in os.walk(dir_path):
        for name in sorted(files):
            file_path = os.path.join(root, name)
            rel = os.path.relpath(file_path, dir_path)
            if name.lower() == "videos.txt":
                all_items.append((file_path, rel, "manifest"))
            elif os.path.splitext(name)[1].lower() in _BATCH_EXTS:
                all_items.append((file_path, rel, "file"))

    total = len(all_items)
    print(f"\n批量入库：目录 {dir_path}，共 {total} 项待处理"
          f"（chunk_size={chunk_size}, chunk_overlap={chunk_overlap}）")

    successes: list[str] = []
    failures: list[tuple[str, str]] = []
    skipped: list[str] = []

    for idx, (file_path, rel, kind) in enumerate(all_items, start=1):
        if kind == "manifest":
            print(f"\n[{idx}/{total}] videos.txt {rel}")
            ok_items, err_items = _process_video_manifest(
                file_path, rel, db_path, db, embedder=embedder,
            )
            successes.extend(ok_items)
            failures.extend(err_items)
            continue
        try:
            visibility, dept_id, owner_user_id = _resolve_visibility(rel)
            print(f"\n[{idx}/{total}] 入库 {rel}（{_scope_tag(visibility, dept_id, owner_user_id)}）")
            process_file(
                file_path, visibility, dept_id, owner_user_id, db_path,
                db=db, chunk_size=chunk_size, chunk_overlap=chunk_overlap,
                verbose=verbose, embedder=embedder,
            )
        except Exception as exc:
            failures.append((rel, str(exc)))
            print(f"入库失败：{rel} -> {exc}")
            continue
        successes.append(rel)

    # 补录未扫描到但在 walk 中被跳过的文件（纯格式不支持）
    for root, _dirs, files in os.walk(dir_path):
        for name in sorted(files):
            rel = os.path.relpath(os.path.join(root, name), dir_path)
            ext = os.path.splitext(name)[1].lower()
            if ext not in _BATCH_EXTS and name.lower() != "videos.txt":
                skipped.append(rel)

    print("\n========== 批量入库汇总 ==========")
    print(f"成功 {len(successes)} 个，失败 {len(failures)} 个，跳过（不支持格式）{len(skipped)} 个")
    for rel, err in failures:
        print(f"  失败：{rel} -> {err}")

def main() -> None:
    parser = argparse.ArgumentParser(description="创建 Milvus Lite 知识库")
    parser.add_argument(
        "file",
        nargs="?",
        help="文档路径（.pdf/.docx/.txt）或图片路径（.jpg/.jpeg/.png/.bmp），视频传 http(s) URL",
    )
    parser.add_argument(
        "--dir",
        dest="dir_path",
        default=None,
        help="批量入库根目录：public/ 为公共，dept/<dept_id>/ 为部门，private/<user_id>/ 为个人",
    )
    parser.add_argument("--visibility", default="public", choices=["public", "dept", "private"])
    parser.add_argument("--dept-id", default=None, help="visibility=dept 时必填")
    parser.add_argument("--owner-user-id", default=None, help="visibility=private 时必填")
    parser.add_argument("--db-path", default=None, help="覆盖 MILVUS_DB_PATH")
    parser.add_argument("--description", default=None, help="图片/视频的说明文本，缺省用文件名")
    parser.add_argument(
        "--chunk-size",
        type=int,
        default=CHUNK_SIZE,
        help=f"切片单块最大字符数（默认 {CHUNK_SIZE}）",
    )
    parser.add_argument(
        "--chunk-overlap",
        type=int,
        default=CHUNK_OVERLAP,
        help=f"切片块间重叠字符数（默认 {CHUNK_OVERLAP}）",
    )
    parser.add_argument(
        "--verbose",
        action="store_true",
        default=False,
        help="打印每个切片的完整文本（调试用）；默认只打印汇总",
    )
    args = parser.parse_args()

    settings = get_settings()
    db_path = args.db_path or settings.milvus_db_path

    # 批量模式：权限标记由目录结构决定，单文件专属参数不适用
    if args.dir_path:
        if (
            args.file
            or args.visibility != "public"
            or args.dept_id
            or args.owner_user_id
            or args.description
        ):
            raise SystemExit(
                "--dir 批量模式下不接受位置参数及 "
                "--visibility/--dept-id/--owner-user-id/--description"
            )
        process_directory(
            args.dir_path, db_path,
            chunk_size=args.chunk_size,
            chunk_overlap=args.chunk_overlap,
            verbose=args.verbose,
        )
    else:
        if not args.file:
            raise SystemExit("必须提供文档路径，或使用 --dir 指定批量入库目录")
        if args.visibility == "dept" and not args.dept_id:
            raise SystemExit("visibility=dept 时必须提供 --dept-id")
        if args.visibility == "private" and not args.owner_user_id:
            raise SystemExit("visibility=private 时必须提供 --owner-user-id")
        process_file(
            args.file,
            args.visibility,
            args.dept_id,
            args.owner_user_id,
            db_path,
            description=args.description,
            chunk_size=args.chunk_size,
            chunk_overlap=args.chunk_overlap,
            verbose=args.verbose,
        )
    print(f"\n知识库构建完成，数据文件: {db_path}")

if __name__ == "__main__":
    main()
